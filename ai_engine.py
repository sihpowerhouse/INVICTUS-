"""SIH Case AI document extraction and Q&A.

Primary provider: Ollama Cloud.
Fallback provider: Gemini API.
Native text extraction is used first for text-bearing files.
Original uploaded bytes are never modified.
"""
from __future__ import annotations

import base64
import io
import json
import os
import re
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any, Callable

from PIL import Image
from dotenv import load_dotenv

load_dotenv(override=False)

OLLAMA_URL = os.getenv("OLLAMA_URL", "https://ollama.com").rstrip("/")
OLLAMA_API_KEY = os.getenv("OLLAMA_API_KEY", "").strip()
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "gemma4:31b")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
REQUEST_TIMEOUT = int(os.getenv("AI_REQUEST_TIMEOUT", "90"))
CHAT_TIMEOUT = int(os.getenv("AI_CHAT_TIMEOUT", "45"))
PDF_DPI = int(os.getenv("AI_PDF_DPI", "120"))


def _post_json(url: str, payload: dict, headers: dict[str, str], timeout: int = REQUEST_TIMEOUT) -> dict:
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", **headers},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code}: {body[:1000]}") from exc
    except Exception as exc:
        raise RuntimeError(str(exc)) from exc


def _native_pages(data: bytes, filename: str) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        text = data.decode("utf-8", errors="replace").strip()
        return [{"page": 1, "text": text, "confidence": 1.0, "source": "native"}] if text else []

    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(parts).strip()
        return [{"page": 1, "text": text, "confidence": 1.0, "source": "native"}] if text else []

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        pages = []
        for no, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            text = "\n".join(texts).strip()
            if text:
                pages.append({"page": no, "text": text, "confidence": 1.0, "source": "native"})
        return pages

    if ext == ".pdf":
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        pages = []
        for no, page in enumerate(doc, 1):
            text = page.get_text("text").strip()
            pages.append({
                "page": no,
                "text": text,
                "confidence": 1.0 if text else None,
                "source": "native" if text else "scan",
            })
        doc.close()
        return pages

    return []


def _render_pdf_pages(data: bytes, page_numbers: list[int]) -> dict[int, bytes]:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    scale = PDF_DPI / 72.0
    matrix = fitz.Matrix(scale, scale)
    out: dict[int, bytes] = {}
    for page_no in page_numbers:
        page = doc.load_page(page_no - 1)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        out[page_no] = pix.tobytes("png")
    doc.close()
    return out


def _image_to_png(data: bytes) -> bytes:
    im = Image.open(io.BytesIO(data)).convert("RGB")
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def _ollama_vision(image_bytes: bytes, page_no: int) -> str:
    if not OLLAMA_API_KEY:
        raise RuntimeError("OLLAMA_API_KEY is not configured.")
    b64 = base64.b64encode(image_bytes).decode("ascii")
    prompt = (
        "Transcribe this document image exactly. Preserve names, dates, numbers, "
        "IDs, headings and table values. Do not summarize, interpret, correct or invent. "
        f"This is page {page_no}. Return only the readable document text."
    )
    payload = {
        "model": OLLAMA_MODEL,
        "messages": [{"role": "user", "content": prompt, "images": [b64]}],
        "stream": False,
        "options": {"temperature": 0},
    }
    result = _post_json(
        f"{OLLAMA_URL}/api/chat",
        payload,
        {"Authorization": f"Bearer {OLLAMA_API_KEY}"},
    )
    text = (((result.get("message") or {}).get("content")) or "").strip()
    if not text:
        raise RuntimeError("Ollama returned no extracted text.")
    return text


def _ollama_text(prompt: str, timeout: int = CHAT_TIMEOUT) -> str:
    if not OLLAMA_API_KEY:
        raise RuntimeError("OLLAMA_API_KEY is not configured.")
    
    print(f"\nOLLAMA REQUEST\nmodel={OLLAMA_MODEL}")
    
    result = _post_json(
        f"{OLLAMA_URL}/api/chat",
        {
            "model": OLLAMA_MODEL,
            "messages": [{"role": "user", "content": prompt}],
            "stream": False,
            "options": {"temperature": 0},
        },
        {"Authorization": f"Bearer {OLLAMA_API_KEY}"},
        timeout=timeout,
    )
    text = (((result.get("message") or {}).get("content")) or "").strip()
    if not text:
        raise RuntimeError("Ollama returned an empty answer.")
    return text


def _gemini_generate(prompt: str, data: bytes | None = None, mime_type: str | None = None, timeout: int = REQUEST_TIMEOUT) -> str:
    if not GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY is not configured.")
    parts: list[dict[str, Any]] = [{"text": prompt}]
    if data is not None and mime_type:
        parts.append({"inline_data": {"mime_type": mime_type, "data": base64.b64encode(data).decode("ascii")}})
    payload = {"contents": [{"parts": parts}]}
    result = _post_json(
        f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent",
        payload,
        {"x-goog-api-key": GEMINI_API_KEY},
        timeout=timeout,
    )
    candidates = result.get("candidates") or []
    if not candidates:
        raise RuntimeError(f"Gemini returned no candidates: {result}")
    out = []
    for part in (candidates[0].get("content") or {}).get("parts") or []:
        if part.get("text"):
            out.append(part["text"])
    text = "\n".join(out).strip()
    if not text:
        raise RuntimeError("Gemini returned an empty answer.")
    return text


def _gemini_document(data: bytes, filename: str, progress_callback: Callable[[str, int | None], None] | None = None) -> list[dict]:
    ext = Path(filename).suffix.lower()
    mime = "application/pdf" if ext == ".pdf" else "image/png"
    if ext != ".pdf":
        data = _image_to_png(data)
    prompt = (
        "Transcribe this document exactly, page by page. Preserve all names, dates, "
        "numbers, IDs, headings and table values. Do not summarize or invent. "
        "Return plain text with page headings like '--- Page 1 ---'."
    )
    if progress_callback:
        progress_callback("gemini_request", None)
    text = _gemini_generate(prompt, data, mime)
    pages = []
    current = None
    for line in text.splitlines():
        m = re.match(r"\s*---\s*Page\s+(\d+)\s*---\s*$", line, re.I)
        if m:
            if current:
                pages.append(current)
            current = {"page": int(m.group(1)), "text": "", "confidence": None, "source": "gemini"}
        elif current is not None:
            current["text"] += ("\n" if current["text"] else "") + line
    if current:
        pages.append(current)
    if not pages:
        pages = [{"page": 1, "text": text, "confidence": None, "source": "gemini"}]
    return pages


def estimate_document_seconds(data: bytes, filename: str) -> int:
    """Return a conservative user-facing estimate; never a guarantee."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            native = _native_pages(data, filename)
            scanned = sum(1 for p in native if not p["text"].strip())
            if scanned == 0:
                return 5
            return min(300, 20 + scanned * 35)
        if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
            return 45
        if ext in {".txt", ".docx", ".pptx"}:
            return 5
    except Exception:
        pass
    return 60


def extract_document(data: bytes, filename: str, progress_callback: Callable[[str, int | None], None] | None = None) -> dict:
    native = _native_pages(data, filename)
    ext = Path(filename).suffix.lower()

    # Digital documents can be handled without an AI request.
    if ext in {".txt", ".docx", ".pptx"}:
        pages = native
        return {"text": "\n\n".join(f"--- Page {p['page']} ---\n{p['text']}" for p in pages),
                "pages": pages, "confidence": 1.0 if pages else None,
                "provider": "native", "model": "native", "fallback_used": False}

    if ext == ".pdf":
        scanned = [p["page"] for p in native if not p["text"].strip()]
        if not scanned:
            return {"text": "\n\n".join(f"--- Page {p['page']} ---\n{p['text']}" for p in native),
                    "pages": native, "confidence": 1.0 if native else None,
                    "provider": "native", "model": "PyMuPDF", "fallback_used": False}
        try:
            images = _render_pdf_pages(data, scanned)
            by_page = {p["page"]: p for p in native}
            total = len(scanned)
            for idx, page_no in enumerate(scanned, 1):
                if progress_callback:
                    progress_callback(f"ollama_page_{page_no}", int(20 + (idx - 1) * 60 / max(total, 1)))
                try:
                    by_page[page_no]["text"] = _ollama_vision(images[page_no], page_no)
                except Exception:
                    if progress_callback:
                        progress_callback("ollama_failed_switching_to_gemini", int(20 + (idx - 1) * 60 / max(total, 1)))
                    raise
                by_page[page_no]["source"] = "ollama"
                by_page[page_no]["confidence"] = None
                if progress_callback:
                    progress_callback(f"ollama_page_{page_no}_complete", int(20 + idx * 60 / max(total, 1)))
            pages = [by_page[n] for n in sorted(by_page)]
            return {"text": "\n\n".join(f"--- Page {p['page']} ---\n{p['text']}" for p in pages if p['text'].strip()),
                    "pages": pages, "confidence": None, "provider": "ollama", "model": OLLAMA_MODEL, "fallback_used": False}
        except Exception:
            pages = _gemini_document(data, filename, progress_callback)
            return {"text": "\n\n".join(f"--- Page {p['page']} ---\n{p['text']}" for p in pages),
                    "pages": pages, "confidence": None, "provider": "gemini", "model": GEMINI_MODEL, "fallback_used": True}

    if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        try:
            if progress_callback:
                progress_callback("ollama_request", 20)
            text = _ollama_vision(_image_to_png(data), 1)
            if progress_callback:
                progress_callback("ollama_complete", 85)
            pages = [{"page": 1, "text": text, "confidence": None, "source": "ollama"}]
            return {"text": text, "pages": pages, "confidence": None,
                    "provider": "ollama", "model": OLLAMA_MODEL, "fallback_used": False}
        except Exception:
            if progress_callback:
                progress_callback("ollama_failed_switching_to_gemini", 25)
            pages = _gemini_document(data, filename, progress_callback)
            return {"text": "\n\n".join(f"--- Page {p['page']} ---\n{p['text']}" for p in pages),
                    "pages": pages, "confidence": None, "provider": "gemini", "model": GEMINI_MODEL, "fallback_used": True}

    raise ValueError("Unsupported AI document type.")


def answer_question(prompt: str) -> dict:
    """Answer with Ollama first, then Gemini fallback, with bounded time per provider."""
    ollama_error = None
    try:
        return {"text": _ollama_text(prompt, CHAT_TIMEOUT), "provider": "ollama", "model": OLLAMA_MODEL, "fallback_used": False}
    except Exception as exc:
        ollama_error = str(exc)
    try:
        return {"text": _gemini_generate(prompt, timeout=CHAT_TIMEOUT), "provider": "gemini", "model": GEMINI_MODEL, "fallback_used": True}
    except Exception as exc:
        raise RuntimeError(f"Ollama failed: {ollama_error}; Gemini fallback failed: {exc}") from exc


def chunk_pages(pages: list[dict], chunk_size: int = 3500) -> list[dict]:
    """Split pages into chunks with a GLOBAL index per immutable version.

    The database constraint is UNIQUE(version_id, chunk_index), so resetting
    chunk_index to zero on every page causes multi-page documents to fail.
    """
    chunks = []
    global_index = 0
    for p in pages:
        text = (p.get("text") or "").strip()
        if not text:
            continue
        for i in range(0, len(text), chunk_size):
            chunks.append({
                "page": p.get("page", 1),
                "chunk_index": global_index,
                "text": text[i:i + chunk_size],
            })
            global_index += 1
    return chunks
